from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app.core.config import settings
from app.main import create_app
from app.schemas.ingestion import ClassifiedEmail
from app.services.ingestion import ingest_email


def email(*names, **kwargs):
    return ClassifiedEmail(email_id="test", tag="SI", attachments=[
        {"filename": Path(name).name, "path": name} for name in names
    ], **kwargs)


def test_text_and_isolated_errors(tmp_path):
    (tmp_path / "SI.txt").write_text("Shipping instruction\nShipper: ABC", encoding="utf-8-sig")
    (tmp_path / "bad.pdf").write_bytes(b"not a PDF")
    result = ingest_email(email("SI.txt", "bad.pdf", "missing.txt", "../secret.txt"), tmp_path)
    assert result.status == "partial"
    assert result.requires_human_review
    assert result.documents[0].document_type == "SI"
    assert "Shipper: ABC" in result.documents[0].text
    assert all(d.status == "error" for d in result.documents[1:])


def test_docx_and_workbook(tmp_path):
    from docx import Document
    from openpyxl import Workbook

    document = Document()
    document.add_paragraph("BILL OF LADING")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Consignee"
    table.cell(0, 1).text = "Example Ltd"
    document.save(tmp_path / "draft.docx")
    workbook = Workbook()
    workbook.active.title = "Shipping"
    workbook.active.append(["Shipping instruction", "Weight", 12500])
    workbook.create_sheet("Containers").append(["Count", 3])
    workbook.save(tmp_path / "SI.xlsx")
    result = ingest_email(email("draft.docx", "SI.xlsx"), tmp_path)
    assert result.status == "ok"
    assert not result.requires_human_review
    assert result.documents[0].document_type == "BL"
    assert "Consignee\tExample Ltd" in result.documents[0].text
    assert result.documents[0].pages is None
    assert {s.sheet_name for s in result.documents[1].segments} == {"Shipping", "Containers"}
    assert "12500" in result.documents[1].text


def test_scanned_pdf_ocr_and_failure(tmp_path):
    from PIL import Image

    image = Image.new("RGB", (400, 200), "white")
    image.save(tmp_path / "SI.pdf", "PDF", save_all=True, append_images=[image])
    calls = []

    def ocr(image):
        calls.append(image.size)
        return "Shipping instruction: ABC"

    result = ingest_email(email("SI.pdf"), tmp_path, ocr=ocr)
    assert len(calls) == 2
    assert result.status == "ok"
    assert result.documents[0].pages == 2
    assert result.documents[0].ocr_used
    assert [s.page_number for s in result.documents[0].segments] == [1, 2]

    def broken_ocr(image):
        raise RuntimeError("missing OCR engine")

    failed = ingest_email(email("SI.pdf"), tmp_path, ocr=broken_ocr)
    assert failed.status == "error"
    assert failed.documents[0].segments[0].error

    calls.clear()

    def partial_ocr(image):
        calls.append(image.size)
        if len(calls) == 2:
            raise RuntimeError("OCR timeout")
        return "Shipping instruction: ABC"

    partial = ingest_email(email("SI.pdf"), tmp_path, ocr=partial_ocr)
    assert partial.status == "partial"
    assert partial.documents[0].status == "partial"
    assert "ABC" in partial.documents[0].text
    assert partial.documents[0].segments[1].error


def test_gate_and_empty_input(tmp_path):
    with pytest.raises(ValidationError):
        ClassifiedEmail(email_id="test", tag="SPAM", attachments=[])
    assert ingest_email(email("missing.txt", requires_human_review=True), tmp_path).status == "review_required"
    assert ingest_email(email(), tmp_path).status == "error"


def test_path_and_size_limits(tmp_path, monkeypatch):
    from app.services import ingestion

    target = tmp_path / "SI.txt"
    target.write_text("Shipping instruction", encoding="utf-8")
    absolute = ingest_email(email(str(target)), tmp_path)
    assert absolute.documents[0].status == "error"
    monkeypatch.setattr(ingestion, "MAX_BYTES", 4)
    oversized = ingest_email(email("SI.txt"), tmp_path)
    assert "limit" in oversized.documents[0].error


def test_api(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "bundle_dir", tmp_path)
    (tmp_path / "attachments").mkdir()
    (tmp_path / "attachments" / "DL.txt").write_text("BILL OF LADING: ABC", encoding="utf-8")
    payload = {"email_id": "test", "tag": "DL", "attachments": [{"filename": "DL.txt", "path": "DL.txt", "document_type": "BL"}]}
    with TestClient(create_app()) as client:
        response = client.post("/api/v1/ingestion", json=payload)
        assert response.status_code == 200
        assert response.json()["documents"][0]["document_type"] == "BL"
        assert response.json()["tag"] == "DL"
        payload["tag"] = "GENERAL"
        assert client.post("/api/v1/ingestion", json=payload).status_code == 422


def test_encodings_empty_unsupported_and_formulas(tmp_path):
    from openpyxl import Workbook

    (tmp_path / "SI.txt").write_text("Shipping instruction: 港口", encoding="utf-16")
    (tmp_path / "empty.txt").write_bytes(b"")
    (tmp_path / "old.doc").write_bytes(b"legacy")
    workbook = Workbook()
    workbook.active.append(["SI", "=1+2"])
    workbook.save(tmp_path / "SI.xlsx")
    result = ingest_email(email("SI.txt", "empty.txt", "old.doc", "SI.xlsx"), tmp_path)
    assert "港口" in result.documents[0].text
    assert result.documents[1].status == "error"
    assert "Unsupported format" in result.documents[2].error
    assert "=1+2" in result.documents[3].text
    assert result.documents[3].warnings


def test_digital_pdf_does_not_call_ocr(tmp_path):
    # A minimal PDF fixture with a real text stream, no extra fixture dependency.
    stream = b"BT /F1 12 Tf 30 100 Td (SHIPPING INSTRUCTION - Shipper ABC Manufacturing - Port Singapore) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    start = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    data += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()
    (tmp_path / "SI.pdf").write_bytes(data)

    def unexpected_ocr(image):
        pytest.fail("Digital PDF should use its text layer")

    result = ingest_email(email("SI.pdf"), tmp_path, ocr=unexpected_ocr)
    assert result.status == "ok"
    assert "ABC Manufacturing" in result.documents[0].text
    assert not result.documents[0].ocr_used
