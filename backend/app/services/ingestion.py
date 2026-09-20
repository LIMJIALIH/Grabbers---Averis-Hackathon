"""Stage 4: local attachment parsing, independent of the email classifier."""

import re
from contextlib import closing
from pathlib import Path
from threading import Lock
from zipfile import ZipFile

from app.schemas.ingestion import (
    ClassifiedEmail, IngestedDocument, IngestionResult, SourceSegment,
)

MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 100
MAX_CELLS = 200_000
PDF_LOCK = Lock()  # PDFium must not be called concurrently from API worker threads.


def ocr_image(image) -> str:
    import pytesseract

    return pytesseract.image_to_string(image, lang="eng", timeout=60).strip()


def parse_pdf(path: Path, result: IngestedDocument, ocr) -> None:
    import pypdfium2 as pdfium

    with PDF_LOCK, pdfium.PdfDocument(path) as pdf:
        if len(pdf) > MAX_PAGES:
            raise ValueError("PDF exceeds the 100-page limit")
        result.pages = len(pdf)
        for index in range(len(pdf)):
            segment = SourceSegment(location=f"page:{index + 1}", page_number=index + 1, text="")
            try:
                with closing(pdf[index]) as page:
                    with closing(page.get_textpage()) as textpage:
                        segment.text = textpage.get_text_range().strip()
                    # Sparse pages can contain a header above a scanned shipping form.
                    if len(re.sub(r"\s", "", segment.text)) < 40:
                        segment.ocr_used = True
                        scale = min(3, 4000 / max(page.get_size()))
                        with closing(page.render(scale=scale)) as bitmap:
                            with bitmap.to_pil() as image:
                                recognized = ocr(image)
                        if recognized.strip():
                            segment.text = recognized.strip()
                        elif not segment.text:
                            segment.error = "No readable text found on page"
            except Exception as exc:
                segment.error = f"Page parsing/OCR failed ({type(exc).__name__})"
            result.segments.append(segment)


def check_office_archive(path: Path) -> None:
    with ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Office archive exceeds the 100 MB expanded-size limit")


def parse_docx(path: Path, result: IngestedDocument) -> None:
    from docx import Document
    from docx.table import Table

    check_office_archive(path)
    document = Document(path)
    for index, block in enumerate(document.iter_inner_content(), 1):
        text = "\n".join("\t".join(cell.text for cell in row.cells) for row in block.rows) if isinstance(block, Table) else block.text
        if text.strip():
            result.segments.append(SourceSegment(location=f"body:block:{index}", text=text))
    for index, section in enumerate(document.sections, 1):
        for name in ("header", "footer"):
            container = getattr(section, name)
            texts = [p.text for p in container.paragraphs]
            texts.extend("\t".join(c.text for c in row.cells) for table in container.tables for row in table.rows)
            if any(texts):
                result.segments.append(SourceSegment(location=f"section:{index}:{name}", text="\n".join(texts)))
    if document.part.element.xpath(".//w:drawing"):
        result.warnings.append("Word drawings/images are not OCR processed; review embedded content")


def parse_xlsx(path: Path, result: IngestedDocument) -> None:
    from openpyxl import load_workbook

    check_office_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        cells = 0
        for sheet in workbook:
            cells += sheet.max_row * sheet.max_column
            if cells > MAX_CELLS:
                raise ValueError("Workbook exceeds the 200,000-cell limit")
            for index, row in enumerate(sheet.iter_rows(), 1):
                if any(cell.data_type == "f" for cell in row):
                    warning = "Formula expressions are preserved, not calculated; review workbook values"
                    if warning not in result.warnings:
                        result.warnings.append(warning)
                values = ["" if cell.value is None else str(cell.value) for cell in row]
                if any(values):
                    result.segments.append(SourceSegment(location=f"sheet:{sheet.title}:row:{index}", sheet_name=sheet.title, text="\t".join(values)))
    finally:
        workbook.close()


def infer_type(filename: str, text: str) -> str:
    evidence = filename + "\n" + text[:4000]
    si = bool(re.search(r"\bshipping\s+instructions?\b|(?<![a-z])si(?![a-z])", evidence, re.I))
    bl = bool(re.search(r"\bbill\s+of\s+lading\b|(?<![a-z])b/?l(?![a-z])", evidence, re.I))
    dl = bool(re.search(r"(?<![a-z])dl(?![a-z])", filename, re.I))
    matches = [label for label, found in (("SI", si), ("BL", bl), ("DL", dl)) if found]
    return matches[0] if len(matches) == 1 else "UNKNOWN"


def ingest_email(email: ClassifiedEmail, attachment_root: Path, *, ocr=ocr_image) -> IngestionResult:
    """Paths are relative to attachment_root; never fetch arbitrary files or URLs."""
    if email.requires_human_review:
        return IngestionResult(email_id=email.email_id, tag=email.tag, status="review_required", requires_human_review=True)
    documents = []
    root = attachment_root.resolve()
    for attachment in email.attachments:
        result = IngestedDocument(filename=attachment.filename, document_type=attachment.document_type or "UNKNOWN")
        documents.append(result)
        try:
            relative = Path(attachment.path)
            path = (root / relative).resolve()
            if relative.is_absolute() or not path.is_relative_to(root) or not path.is_file():
                raise ValueError("Attachment is unavailable or outside the attachment directory")
            if path.stat().st_size > MAX_BYTES:
                raise ValueError("Attachment exceeds the 25 MB limit")
            extension = path.suffix.lower()
            if extension != Path(attachment.filename).suffix.lower():
                raise ValueError("Filename and stored attachment extensions differ")
            if extension == ".pdf":
                parse_pdf(path, result, ocr)
            elif extension == ".docx":
                parse_docx(path, result)
            elif extension == ".xlsx":
                parse_xlsx(path, result)
            elif extension == ".txt":
                raw = path.read_bytes()
                encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
                result.segments.append(SourceSegment(location="text", text=raw.decode(encoding)))
            else:
                raise ValueError("Unsupported format; export as PDF, DOCX, XLSX, or UTF-8 TXT")
            result.text = "\n\n".join(segment.text for segment in result.segments if segment.text.strip())
            result.ocr_used = any(segment.ocr_used for segment in result.segments)
            if attachment.document_type is None:
                result.document_type = infer_type(attachment.filename, result.text)
            if result.document_type == "UNKNOWN":
                result.warnings.append("Document type is ambiguous; supply an attachment document_type or review")
            if not result.text.strip():
                result.error = "No readable text extracted"
            else:
                result.status = "partial" if any(segment.error for segment in result.segments) else "ok"
        except Exception as exc:
            result.error = str(exc) if isinstance(exc, (ValueError, UnicodeError)) else f"Attachment parsing failed ({type(exc).__name__})"
    successful = sum(document.status == "ok" for document in documents)
    status = "ok" if documents and successful == len(documents) else "partial" if any(d.text for d in documents) else "error"
    review = status != "ok" or any(d.warnings for d in documents)
    return IngestionResult(email_id=email.email_id, tag=email.tag, status=status, requires_human_review=review, documents=documents)
